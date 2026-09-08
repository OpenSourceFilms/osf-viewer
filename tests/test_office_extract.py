"""Regression tests for office_extract.py -- semantic extraction for
DOCX/XLSX/PPTX/ODT/ODS/ODP and legacy DOC/XLS/PPT/RTF (via LibreOffice
headless conversion into the same modern-format extractors), plus
/preview.pdf rendering and its content-addressed cache.

All fixtures are the hand-built files under tests/fixtures/ (see DOCS.md).
Assertions check actual known fixture content, not just "non-empty", so a
regression that silently drops a heading/table/sheet/slide/notes would be
caught, not just a total-failure regression.

Run with: /workspace/services/osf_viewer/.venv/bin/python -m unittest
          tests.test_office_extract -v
(from /workspace/services/osf_viewer/)

Note: several tests invoke real LibreOffice headless conversions (legacy
formats, /preview.pdf) and are correspondingly slower (~4-6s each) -- this
is deliberate, matching the "verify by function" rule for this workspace
rather than mocking soffice out of the legacy-format path entirely.
"""
from __future__ import annotations

import io
import sys
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import cache as _cache  # noqa: E402
import inspectors  # noqa: E402
import office_extract as oe  # noqa: E402

FIXTURES = Path(__file__).resolve().parent / "fixtures"


class DocxTests(unittest.TestCase):
    def setUp(self):
        self.txt = oe.get_text(FIXTURES / "sample.docx", None, "sample.docx", "docx")

    def test_heading_and_title_detected(self):
        self.assertIn("# Main Title", self.txt)
        self.assertIn("# Section One", self.txt)

    def test_paragraph_order_preserved(self):
        title_i = self.txt.index("Main Title")
        section_i = self.txt.index("Section One")
        para_i = self.txt.index("This is a plain paragraph")
        table_heading_i = self.txt.index("A Table Section")
        closing_i = self.txt.index("Closing paragraph")
        self.assertLess(title_i, section_i)
        self.assertLess(section_i, para_i)
        self.assertLess(para_i, table_heading_i)
        self.assertLess(table_heading_i, closing_i)

    def test_table_rendered_intelligibly(self):
        self.assertIn("[table]", self.txt)
        self.assertIn("Name | Value | Notes", self.txt)
        self.assertIn("alpha | 1 | first row", self.txt)

    def test_unicode_preserved(self):
        self.assertIn("café, naïve, 日本語, emoji 🎬", self.txt)


class XlsxTests(unittest.TestCase):
    def setUp(self):
        self.txt = oe.get_text(FIXTURES / "sample.xlsx", None, "sample.xlsx", "xlsx")

    def test_worksheet_names_and_order(self):
        header_line = next(l for l in self.txt.splitlines() if l.startswith("worksheets:"))
        self.assertIn("Sheet_Numbers", header_line)
        self.assertIn("Sheet_Unicode", header_line)
        self.assertLess(header_line.index("Sheet_Numbers"), header_line.index("Sheet_Unicode"))
        self.assertLess(self.txt.index("--- sheet: Sheet_Numbers"),
                         self.txt.index("--- sheet: Sheet_Unicode"))

    def test_bounded_cell_extraction(self):
        self.assertIn("id | name | score", self.txt)
        self.assertIn("1 | alpha | 9.5", self.txt)
        self.assertIn("2 | beta café | 7.25", self.txt)

    def test_unicode_preserved(self):
        self.assertIn("greeting | lang", self.txt)
        self.assertIn("こんにちは | ja", self.txt)
        self.assertIn("Grüße | de", self.txt)


class PptxTests(unittest.TestCase):
    def setUp(self):
        self.txt = oe.get_text(FIXTURES / "sample.pptx", None, "sample.pptx", "pptx")

    def test_slide_boundaries_and_order(self):
        i1 = self.txt.index("--- slide 1 ---")
        i2 = self.txt.index("--- slide 2 ---")
        i3 = self.txt.index("--- slide 3 ---")
        self.assertLess(i1, i2)
        self.assertLess(i2, i3)

    def test_visible_text_extracted(self):
        for i in (1, 2, 3):
            self.assertIn(f"Slide {i} Title", self.txt)
            self.assertIn(f"unicode: café {i}", self.txt)

    def test_speaker_notes_extracted(self):
        for i in (1, 2, 3):
            self.assertIn(f"[notes] Speaker notes for slide {i}.", self.txt)

    def test_unicode_preserved(self):
        self.assertIn("café", self.txt)


class OdfTests(unittest.TestCase):
    def test_odt_heading_paragraph_unicode(self):
        txt = oe.get_text(FIXTURES / "sample.odt", None, "sample.odt", "odt")
        self.assertIn("# ODT Fixture Heading", txt)
        self.assertIn("An ODT paragraph with unicode café.", txt)
        self.assertLess(txt.index("ODT Fixture Heading"), txt.index("An ODT paragraph"))

    def test_ods_sheet_name_and_cells(self):
        txt = oe.get_text(FIXTURES / "sample.ods", None, "sample.ods", "ods")
        self.assertIn("worksheets: 1 -> Sheet1", txt)
        self.assertIn("id | name", txt)
        self.assertIn("1 | café", txt)

    def test_odp_slide_order_and_text(self):
        txt = oe.get_text(FIXTURES / "sample.odp", None, "sample.odp", "odp")
        i1 = txt.index("--- slide 1 ---")
        i2 = txt.index("--- slide 2 ---")
        self.assertLess(i1, i2)
        self.assertIn("ODP fixture slide 1 text cafe", txt)
        self.assertIn("ODP fixture slide 2 text cafe", txt)


class LegacyOfficeConversionTests(unittest.TestCase):
    """The .doc/.xls/.ppt/.rtf fixtures were generated from the same source
    content as their modern sample.docx/xlsx/pptx counterparts (see DOCS.md),
    so the LibreOffice-conversion path must reproduce the same structural
    facts the native modern-format extractor produces -- this is exactly
    what was previously only checked by hand.
    """

    def test_doc_matches_docx_structure(self):
        txt = oe.get_text(FIXTURES / "sample.doc", None, "sample.doc", "doc")
        self.assertIn("[converted from legacy .doc via LibreOffice headless]", txt)
        self.assertIn("# Main Title", txt)
        self.assertIn("# Section One", txt)
        self.assertIn("café, naïve, 日本語, emoji 🎬", txt)
        self.assertIn("[table]", txt)
        self.assertIn("alpha | 1 | first row", txt)

    def test_xls_matches_xlsx_structure(self):
        txt = oe.get_text(FIXTURES / "sample.xls", None, "sample.xls", "xls")
        self.assertIn("[converted from legacy .xls via LibreOffice headless]", txt)
        self.assertIn("Sheet_Numbers", txt)
        self.assertIn("Sheet_Unicode", txt)
        self.assertIn("1 | alpha | 9.5", txt)
        self.assertIn("こんにちは | ja", txt)

    def test_ppt_matches_pptx_structure(self):
        txt = oe.get_text(FIXTURES / "sample.ppt", None, "sample.ppt", "ppt")
        self.assertIn("[converted from legacy .ppt via LibreOffice headless]", txt)
        i1 = txt.index("--- slide 1 ---")
        i2 = txt.index("--- slide 2 ---")
        i3 = txt.index("--- slide 3 ---")
        self.assertLess(i1, i2)
        self.assertLess(i2, i3)
        self.assertIn("[notes] Speaker notes for slide 1.", txt)

    def test_rtf_matches_docx_structure(self):
        txt = oe.get_text(FIXTURES / "sample.rtf", None, "sample.rtf", "rtf")
        self.assertIn("[converted from legacy .rtf via LibreOffice headless]", txt)
        self.assertIn("# Main Title", txt)
        self.assertIn("café, naïve, 日本語, emoji 🎬", txt)


class CorruptFileFallbackTests(unittest.TestCase):
    """The overriding contract (inspectors.get_text) is that an existing
    file NEVER produces an HTTP 500 -- these exercise that through the real
    top-level dispatch (inspectors.get_text), not just office_extract
    directly, since the demotion-to-fallback wrapper lives one layer up.
    """

    def test_corrupt_docx_degrades_cleanly_no_raise(self):
        txt = inspectors.get_text(FIXTURES / "corrupt.docx", None, "corrupt.docx")
        self.assertIsInstance(txt, str)
        self.assertTrue(txt.strip())
        # demoted to the generic binary report -- python-docx couldn't open it
        self.assertIn("handler failed", txt)
        self.assertIn("mime (libmagic)", txt)
        self.assertIn("sha256", txt)

    def test_corrupt_xls_degrades_cleanly_no_raise(self):
        # LibreOffice is tolerant enough to "import" this garbage as a
        # single-cell text sheet rather than failing outright -- either
        # outcome (clean fallback report, or a degenerate-but-well-formed
        # extraction) satisfies the "never 500, always something useful"
        # contract, so this only asserts the no-crash / non-empty part.
        txt = inspectors.get_text(FIXTURES / "corrupt.xls", None, "corrupt.xls")
        self.assertIsInstance(txt, str)
        self.assertTrue(txt.strip())

    def test_corrupt_docx_direct_office_extract_raises_not_crashes_caller(self):
        # office_extract itself is allowed to raise (that's its documented
        # contract -- inspectors.get_text is the safety net) -- confirm it
        # actually does raise here, so the fallback path above is being
        # exercised for a real reason and not accidentally succeeding.
        with self.assertRaises(Exception):
            oe.get_text(FIXTURES / "corrupt.docx", None, "corrupt.docx", "docx")


class TruncationCapTests(unittest.TestCase):
    """Constructs small in-memory workbooks/decks/docs that exceed the
    row/slide caps so the truncation markers are actually exercised, rather
    than relying on a fixture large enough to trigger them.
    """

    def test_xlsx_row_cap_truncates(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Big"
        for r in range(1, oe.MAX_XLSX_ROWS_PER_SHEET + 100):
            ws.append([r, f"row{r}"])
        buf = io.BytesIO()
        wb.save(buf)
        txt = oe.get_text(None, buf.getvalue(), "big.xlsx", "xlsx")
        self.assertIn(f"truncated at {oe.MAX_XLSX_ROWS_PER_SHEET} rows", txt)
        self.assertNotIn(f"row{oe.MAX_XLSX_ROWS_PER_SHEET + 50}", txt)

    def test_pptx_slide_cap_truncates(self):
        import pptx

        prs = pptx.Presentation()
        blank = prs.slide_layouts[6]
        for i in range(oe.MAX_SLIDES + 5):
            slide = prs.slides.add_slide(blank)
            tb = slide.shapes.add_textbox(0, 0, 100, 100)
            tb.text_frame.text = f"slide {i}"
        buf = io.BytesIO()
        prs.save(buf)
        txt = oe.get_text(None, buf.getvalue(), "big.pptx", "pptx")
        self.assertIn(f"truncated at {oe.MAX_SLIDES} slides", txt)

    def test_docx_table_row_cap_truncates(self):
        import docx

        d = docx.Document()
        table = d.add_table(rows=1, cols=1)
        table.rows[0].cells[0].text = "header"
        for i in range(oe.MAX_TABLE_ROWS + 20):
            row = table.add_row()
            row.cells[0].text = f"r{i}"
        buf = io.BytesIO()
        d.save(buf)
        txt = oe.get_text(None, buf.getvalue(), "big.docx", "docx")
        self.assertIn(f"table truncated at {oe.MAX_TABLE_ROWS} rows", txt)


class PreviewPdfCacheTests(unittest.TestCase):
    """Uses a throwaway relpath/mtime cache key (not a real viewer path) so
    this cleans up after itself without disturbing real cached previews.
    """

    def setUp(self):
        self.relpath = "__test_office_extract_cache_probe__/sample.pptx"
        self.data = (FIXTURES / "sample.pptx").read_bytes()
        self.size = len(self.data)
        self.mtime = 1_700_000_000.0
        self._key = _cache._key_path(self.relpath, self.size, self.mtime, ".pdf")
        if self._key.exists():
            self._key.unlink()

    def tearDown(self):
        if self._key.exists():
            self._key.unlink()

    def test_repeat_preview_request_hits_cache_not_soffice(self):
        real_convert = oe._convert_with_soffice
        calls = []

        def counting(*args, **kwargs):
            calls.append(1)
            return real_convert(*args, **kwargs)

        with mock.patch.object(oe, "_convert_with_soffice", side_effect=counting):
            out1 = oe.get_preview_pdf(None, self.data, "sample.pptx", "pptx",
                                       self.relpath, self.size, self.mtime)
            out2 = oe.get_preview_pdf(None, self.data, "sample.pptx", "pptx",
                                       self.relpath, self.size, self.mtime)

        self.assertEqual(len(calls), 1, "second request should be served from cache, "
                                         "not re-invoke LibreOffice")
        self.assertTrue(out1.startswith(b"%PDF-"))
        self.assertEqual(out1, out2)

    def test_first_request_actually_produces_valid_pdf(self):
        out = oe.get_preview_pdf(None, self.data, "sample.pptx", "pptx",
                                  self.relpath, self.size, self.mtime)
        self.assertIsNotNone(out)
        self.assertTrue(out.startswith(b"%PDF-"))


if __name__ == "__main__":
    unittest.main()
