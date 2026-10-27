#!/usr/bin/env python3
"""Generate disposable Office-fixture replacement candidates from source.

Visible content is authored here. No existing fixture or third-party media is
copied. Office libraries still provide their own package infrastructure under
separately licensed dependencies: this is not legal provenance certification.

Never writes into tests/fixtures/ and refuses an existing destination.
Optional legacy conversion uses a disposable LibreOffice user profile.
"""
from __future__ import annotations

import argparse
from pathlib import Path
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET


EXCLUDED_COMPONENTS = (
    "docProps/thumbnail.jpeg", "docProps/thumbnail.png",
    "customXml/", "ppt/printerSettings/",
)


def sanitize_generated_package(file: Path) -> None:
    """Strip template extras and their references from freshly created files."""
    tmp = file.with_name(file.name + ".sanitizing")
    try:
        with zipfile.ZipFile(file, "r") as src, zipfile.ZipFile(tmp, "w") as dst:
            for info in src.infolist():
                name = info.filename
                if any(name.startswith(prefix) for prefix in EXCLUDED_COMPONENTS):
                    continue
                content = src.read(info)
                if name.endswith(".rels") or name == "[Content_Types].xml":
                    root = ET.fromstring(content)
                    if root.tag.startswith("{"):
                        ET.register_namespace("", root.tag[1:].split("}", 1)[0])
                    for entry in list(root):
                        target = entry.get("Target", "") or entry.get("PartName", "")
                        if ("thumbnail.jpeg" in target or "thumbnail.png" in target
                                or "customXml/" in target
                                or "printerSettings/" in target):
                            root.remove(entry)
                    content = ET.tostring(
                        root, encoding="utf-8", xml_declaration=True
                    )
                dst.writestr(name, content, compress_type=zipfile.ZIP_DEFLATED)
        tmp.replace(file)
    finally:
        if tmp.exists():
            tmp.unlink()


def create_modern(out: Path) -> None:
    import docx
    import openpyxl
    import pptx
    from pptx.util import Inches
    from odf.opendocument import (
        OpenDocumentText, OpenDocumentSpreadsheet, OpenDocumentPresentation,
    )
    from odf import text, table, draw

    document = docx.Document()
    document.core_properties.title = "OSF synthetic document fixture"
    document.core_properties.author = "Open Source Films"
    document.add_heading("Main Title", 0)
    document.add_heading("Section One", level=1)
    document.add_paragraph(
        "This is a plain paragraph with café, naïve, 日本語, emoji 🎬."
    )
    document.add_heading("A Table Section", level=2)
    t = document.add_table(rows=1, cols=3)
    for cell, value in zip(
        t.rows[0].cells, ("Name", "Value", "Notes")
    ):
        cell.text = value
    row = t.add_row()
    for cell, value in zip(row.cells, ("alpha", "1", "first row")):
        cell.text = value
    document.add_paragraph("Closing paragraph.")
    document.save(out / "sample.docx")

    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "Sheet_Numbers"
    a.append(("id", "name", "score"))
    a.append((1, "alpha", 9.5))
    a.append((2, "beta café", 7.25))
    b = wb.create_sheet("Sheet_Unicode")
    b.append(("greeting", "lang"))
    b.append(("こんにちは", "ja"))
    b.append(("Grüße", "de"))
    wb.save(out / "sample.xlsx")

    prs = pptx.Presentation()
    for i in range(1, 4):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(2)
        ).text_frame
        frame.text = f"Slide {i} Title\nunicode: café {i}"
        slide.notes_slide.notes_text_frame.text = (
            f"Speaker notes for slide {i}."
        )
    prs.save(out / "sample.pptx")
    sanitize_generated_package(out / "sample.docx")
    sanitize_generated_package(out / "sample.pptx")

    odt = OpenDocumentText()
    odt.text.addElement(text.H(outlinelevel=1, text="ODT Fixture Heading"))
    odt.text.addElement(
        text.P(text="An ODT paragraph with unicode café.")
    )
    odt.save(str(out / "sample.odt"))

    ods = OpenDocumentSpreadsheet()
    sheet = table.Table(name="Sheet1")
    for row in [("id", "name"), ("1", "café")]:
        tr = table.TableRow()
        for value in row:
            cell = table.TableCell(valuetype="string")
            cell.addElement(text.P(text=value))
            tr.addElement(cell)
        sheet.addElement(tr)
    ods.spreadsheet.addElement(sheet)
    ods.save(str(out / "sample.ods"))

    odp = OpenDocumentPresentation()
    for i in (1, 2):
        page = draw.Page(name=f"Slide {i}", masterpagename="Default")
        frame = draw.Frame(
            width="20cm", height="5cm", x="1cm", y="1cm"
        )
        box = draw.TextBox()
        box.addElement(text.P(text=f"ODP fixture slide {i} text cafe"))
        frame.addElement(box)
        page.addElement(frame)
        odp.presentation.addElement(page)
    odp.save(str(out / "sample.odp"))

    # Intentionally invalid inputs written from literals, not old binaries.
    (out / "corrupt.docx").write_bytes(
        b"PK\x03\x04OSF synthetic invalid DOCX input"
    )
    (out / "corrupt.xls").write_bytes(
        b"OSF synthetic invalid XLS input"
    )


def assert_clean_packages(out: Path) -> None:
    """Reject inherited template extras and external OOXML relationships."""
    for name in (
        "sample.docx", "sample.xlsx", "sample.pptx",
        "sample.odt", "sample.ods", "sample.odp",
    ):
        with zipfile.ZipFile(out / name) as archive:
            for item in archive.namelist():
                if (
                    any(item.startswith(p) for p in EXCLUDED_COMPONENTS)
                    or item.startswith("/")
                    or ".." in Path(item).parts
                ):
                    raise RuntimeError(
                        f"Unexpected package component in {name}: {item}"
                    )
                if item.endswith(".rels"):
                    root = ET.fromstring(archive.read(item))
                    if any(
                        el.get("TargetMode", "").lower() == "external"
                        for el in root
                    ):
                        raise RuntimeError(
                            f"External relationship in {name}: {item}"
                        )


def create_legacy(out: Path) -> None:
    specs = [
        ("sample.docx", "doc:MS Word 97", "sample.doc"),
        ("sample.xlsx", "xls:MS Excel 97", "sample.xls"),
        ("sample.pptx", "ppt:MS PowerPoint 97", "sample.ppt"),
        ("sample.docx", "rtf:Rich Text Format", "sample.rtf"),
    ]
    with tempfile.TemporaryDirectory(
        prefix="osf_soffice_profile_"
    ) as profile:
        for source, fmt, target in specs:
            cmd = [
                "libreoffice",
                "-env:UserInstallation=" + Path(profile).as_uri(),
                "--headless", "--convert-to", fmt,
                "--outdir", str(out), str(out / source),
            ]
            result = subprocess.run(
                cmd, capture_output=True, text=True,
                timeout=75, check=False,
            )
            if result.returncode or not (out / target).is_file():
                raise RuntimeError(
                    f"Could not generate {target}: "
                    f"{result.stderr[-300:]} {result.stdout[-300:]}"
                )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out", type=Path, required=True,
        help="New directory outside the source fixture tree",
    )
    parser.add_argument(
        "--with-legacy", action="store_true",
        help="Also create Office 97 formats using local LibreOffice",
    )
    args = parser.parse_args()
    out = args.out.resolve()
    if "tests/fixtures" in out.as_posix() or out.exists():
        parser.error("output must be a new directory, never tests/fixtures")
    out.mkdir(parents=False, exist_ok=False)
    create_modern(out)
    assert_clean_packages(out)
    if args.with_legacy:
        create_legacy(out)
    for item in sorted(out.iterdir()):
        print(f"{item.name}: {item.stat().st_size} bytes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
